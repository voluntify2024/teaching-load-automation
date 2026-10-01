import os
import re
import math
import unicodedata
from collections import Counter, OrderedDict, defaultdict
from copy import copy
from datetime import date, datetime, timedelta
from difflib import SequenceMatcher
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# ---------------------------------------------------------------------------
# Settings
# ---------------------------------------------------------------------------
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
INPUT_FILE = os.path.join(
    BASE_DIR, 'input_files',
    'Kevadsemester 2026_27, Programmijuhtide poolt sisu semestri planeerimiseks_Yuliya Nevar, IT.xlsx')
OUTPUT_DIR = os.path.join(BASE_DIR, 'output_files')

GREY = 'D9D9D9'  # "the teacher has no such group"
SHORT_DAY = 'E2EFDA'  # colour of "Lühendatud tööpäev" (it is used in the calendar)
HEADER_BLUE = '0070C0'  # header of the subject table (as in the sample)
FONT = 'Times New Roman'

# Optional text for the block "Olulised tähtajad" (sample: cells Z2:AG8).
# The shared workbook has no such text, so it is empty by default.
DEADLINES_TEXT = ''

# Manual name merging, only if the automatic merging fails (see the report).
# {'name as written in the table': 'canonical name'}
NAME_OVERRIDES = {}

# Text that is not a group name
PLACEHOLDERS = ('kellegi', '???', '?')

# Sheet columns are found by the BEGINNING of the header text (normalised),
# because the column order differs between the specialty sheets.
HEADER_KEYS = OrderedDict([
    ('group', lambda h: h.startswith('õpperühm')),
    ('others', lambda h: h.startswith('kelle tunniplaaniga')),
    ('module', lambda h: h.startswith('õppeaasta')),
    ('code', lambda h: h.startswith('aine kood')),
    ('name', lambda h: h.startswith('aine nimetus')),
    ('eap', lambda h: h.startswith('aine maht')),
    ('subgrp', lambda h: h.startswith('gruppid õppeaines')),
    ('pairs', lambda h: h.startswith('kontaktpaare')),
    ('ngroups', lambda h: h == 'grupid'),
    ('type', lambda h: h.startswith('kohustuslik')),
    ('lead', lambda h: h.startswith('vastutav')),
    ('co', lambda h: h.startswith('õppejõud kes on veel')),
    ('status', lambda h: h.startswith('lepinguline')),
    ('contact', lambda h: h.startswith('uue õppejõu kontaktid')),
    ('notes', lambda h: h in ('märkused', 'kommentaarid')),
    ('rgroup', lambda h: h == 'rühm'),
])
DAY_LABELS = ['E', 'T', 'K', 'N', 'R', 'L', 'P']  # Mon ... Sun in the calendar


# ---------------------------------------------------------------------------
# Text helpers
# ---------------------------------------------------------------------------
def clean(value):
    """Cell value -> clean string (NFC, no NBSP, single spaces)."""
    if value is None:
        return ''
    text = unicodedata.normalize('NFC', str(value)).replace('\xa0', ' ')
    return ' '.join(text.split())


def normalize(text):
    return clean(text).casefold()


def to_number(value):
    """'6' -> 6, 6.0 -> 6, '10-12' stays text, empty stays ''."""
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str) and re.fullmatch(r'\d+', value.strip()):
        return int(value.strip())
    return value if value is not None else ''


def fill(rgb):
    return PatternFill(start_color=rgb, end_color=rgb, fill_type='solid')


def safe_filename(text):
    """ASCII-only file name (Aljona Kordontśuk -> Aljona Kordontsuk): Windows ZIP, Teams and
    e-mail often fail on letters like š, ž, õ, ś. The real name stays in cell A1 of the file."""
    text = unicodedata.normalize('NFKD', text).encode('ascii', 'ignore').decode('ascii')
    return re.sub(r'[\\/*?:"<>|]+', '', text).strip()


# ---------------------------------------------------------------------------
# People: names, markers (TT / L), name merging
# ---------------------------------------------------------------------------
def looks_like_name(text):
    """A person: >= 2 words, each starts with a capital, no digits.
    'Teevad 2 rühma' is a comment, not a person."""
    words = [w for w in re.split(r'[\s\-]+', text) if w]
    return (len(words) >= 2 and not re.search(r'\d', text)
            and all(w[0].isupper() or w.casefold() in ('von', 'de', 'van', 'da') for w in words))


def split_people(cell):
    """Splits one cell into people.
    'Ülle Säälik TT, Andu Rämmer (1 paar)' -> [(name, marker, note), ...]
    marker: 'TT' / 'lepinguline' / None (written next to the name).
    Text that is not a name is returned as (None, None, text) = a free note."""
    text = clean(cell)
    if not text or text in ('?', '-', '???'):
        return []
    result = []
    for part in re.split(r'\s*[,;&+]\s*|\s+ja\s+', text):
        part = part.strip()
        if not part:
            continue
        note = ''
        m = re.search(r'\(([^)]*)\)', part)
        if m:
            note = m.group(1).strip()
            part = (part[:m.start()] + part[m.end():]).strip()
        marker = None
        if note.upper() in ('TT', 'L'):  # 'Viivika Roostar (TT)'
            marker, note = ('TT' if note.upper() == 'TT' else 'lepinguline'), ''
        m = re.search(r'\s+(TT|L)\.?$', part)  # 'Ülle Säälik TT'
        if m:
            marker = 'TT' if m.group(1) == 'TT' else 'lepinguline'
            part = part[:m.start()].strip()
        if len(part) < 3:
            continue
        if not looks_like_name(part):
            result.append((None, None, clean(part + (' (' + note + ')' if note else ''))))
        else:
            result.append((part, marker, note))
    return result


def status_tokens(cell):
    """'L, L' -> [('lepinguline', False), ('lepinguline', False)];
    'lepinguline/ tunnitasulised' -> [('lepinguline', False), ('TT', True)]
    (second item = plural form, it covers several people). 'tunnitasuline' = TT."""
    tokens = []
    for t in re.split(r'[,/;]|\s+ja\s+', normalize(cell)):
        t = t.strip(' .')
        if not t or t == '?':
            continue
        plural = bool(re.search(r'(lised|ulised)$', t))
        if t.startswith('tt') or t.startswith('tunnitas'):
            tokens.append(('TT', plural))
        elif t == 'l' or t.startswith('leping'):
            tokens.append(('lepinguline', plural))
    return tokens


def name_tokens(name):
    return [t for t in re.split(r'[\s\-–]+', normalize(name)) if t]


def same_person(a, b):
    """Same first name AND the surname (last word) of one name occurs in the other.
    'Kayleigh Mary Kleiva' ~ 'Kayleigh Kleiva';  'Merju Lomp' ~ 'Merju Lomp-Alev';
    'Anne Mari Tamm' !~ 'Anne Mari Kuusk' (only the middle name is shared)."""
    ta, tb = name_tokens(a), name_tokens(b)
    if len(ta) < 2 or len(tb) < 2:
        return ta == tb
    if ta[0] != tb[0]:
        return False
    return ta[-1] in tb[1:] or tb[-1] in ta[1:]


class NameResolver:
    """Maps every spelling of a name to ONE key (so a teacher gets ONE file)."""

    def __init__(self, raw_names):
        self.key_of = {}
        representatives = []  # [(key, name)]
        for raw in raw_names:
            name = clean(raw)
            if name in self.key_of:
                continue
            if name in NAME_OVERRIDES:
                name_for_key = NAME_OVERRIDES[name]
            else:
                name_for_key = name
            hit = next((k for k, rep in representatives if same_person(name_for_key, rep)), None)
            if hit is None:
                hit = normalize(name_for_key)
                representatives.append((hit, name_for_key))
            self.key_of[name] = hit

    def __call__(self, name):
        return self.key_of[clean(name)]


# ---------------------------------------------------------------------------
# Groups
# ---------------------------------------------------------------------------
def parse_groups(text, default_program, programs):
    """'IT1, IT2' -> [('IT',1),('IT',2)];  'EDL 2, EDL 3 + IT' -> [('EDL',2),('EDL',3),('IT',None)]
    'KMK IK1 2025/26' -> [('KMK',1)];  unknown code ('KK 2') -> the programme of the sheet."""
    groups = []
    for piece in re.split(r'\s*[,;+]\s*|\s+ja\s+', clean(text)):
        if not piece or piece.casefold().startswith(PLACEHOLDERS) or '?' in piece:
            continue
        m = re.match(r'([A-ZÕÄÖÜ]{2,5})', piece)
        if not m:
            continue
        program = m.group(1) if m.group(1) in programs else default_program
        y = re.search(r'(?<!\d)(\d)(?!\d)', piece)
        groups.append((program, int(y.group(1)) if y else None))
    return groups


# ---------------------------------------------------------------------------
# Block 1: reading the shared workbook
# ---------------------------------------------------------------------------
def find_special_sheets(wb):
    """The calendar sheet (name contains 'põhi'), 'Abiinfo' and the specialty sheets."""
    cal = next(ws for ws in wb if 'põhi' in ws.title.casefold())
    abi = next(ws for ws in wb if ws.title.casefold() == 'abiinfo')
    programs = [ws for ws in wb if ws is not cal and ws is not abi]
    return cal, abi, programs


def read_study_modes(abi):
    """Abiinfo table: 'EDL' -> 'sessioonõpe', 'IT' -> 'päevaõpe' ..."""
    modes = {}
    for row in abi.iter_rows():
        code, kind = clean(row[1].value), normalize(row[4].value)
        if re.fullmatch(r'[A-ZÕÄÖÜ]{2,5}', code) and kind.endswith('õpe'):
            modes[code] = kind
    return modes


def map_header(cells):
    cols = {}
    for c in cells:
        h = normalize(c.value)
        if not h:
            continue
        for key, test in HEADER_KEYS.items():
            if key not in cols and test(h):
                cols[key] = c.column - 1
                break
    return cols


def read_program_sheet(ws, programs):
    """One record per subject row. A sheet may contain several header rows,
    every header row re-maps the columns."""
    program = re.sub(r'\s+', '', ws.title).upper()
    records, cols = [], None
    for row in ws.iter_rows():
        vals = [clean(c.value) for c in row]
        joined = ' '.join(v.casefold() for v in vals)
        if 'aine kood' in joined and 'vastutav' in joined:
            cols = map_header(row)
            continue
        if cols is None:
            continue

        def get(key):
            i = cols.get(key)
            return vals[i] if i is not None and i < len(vals) else ''

        code, name = get('code'), get('name')
        if not (code or name) or name.upper().startswith('KOKKU'):
            continue
        lead_all, co_all = split_people(get('lead')), split_people(get('co'))
        rec = dict(
            sheet=ws.title, row=row[0].row, program=program,
            group_primary=get('group'), group_other=get('others'), module=get('module'),
            code=code, name=name, eap=to_number(get('eap')), pairs=to_number(get('pairs')),
            subgrp=get('subgrp'), ngroups=get('ngroups'), type=get('type'),
            notes=get('notes'), rgroup=get('rgroup'), contact=get('contact'),
            status_raw=get('status'),
            leads=[p for p in lead_all if p[0]], cos=[p for p in co_all if p[0]],
            free_notes=[p[2] for p in lead_all + co_all if not p[0]],
        )
        tokens = parse_groups(rec['group_primary'], program, programs)
        for t in parse_groups(rec['group_other'], program, programs):
            if t not in tokens:
                tokens.append(t)
        rec['tokens'] = tokens
        records.append(rec)
    return records


def assign_row_status(rec, people, warnings):
    """The status cell describes the ROW, not one person -> split it between the people:
      1) as many statuses as people          -> one by one;
      2) as many statuses as responsible     -> only the responsible ones;
      3) one status                          -> the responsible (plural 'lepingulised' -> everybody);
      4) two statuses, 3+ people             -> first / the rest (reported as 'derived')."""
    tokens = status_tokens(rec['status_raw'])
    if not tokens:
        return {}
    names = [p[0][0] for p in people]
    leads = [p[0] for p in rec['leads']]
    if len(tokens) == len(names):
        return {n: t for n, (t, _) in zip(names, tokens)}
    if len(tokens) == len(leads):
        return {n: t for n, (t, _) in zip(leads, tokens)}
    if len(tokens) == 1:
        t, plural = tokens[0]
        return {n: t for n in (names if plural else leads)}
    if len(tokens) == 2 and len(names) >= 3:
        (t1, p1), (t2, p2) = tokens
        if p1:
            res = {n: t1 for n in names[:-1]}
            res[names[-1]] = t2
        else:
            res = {names[0]: t1}
            res.update({n: t2 for n in names[1:]})
        warnings.append(('Staatus tuletatud (kontrolli)', f"{rec['sheet']} r.{rec['row']}",
                         f"'{rec['status_raw']}' -> {res}"))
        return res
    warnings.append(('Staatus ei sobitu nimedega', f"{rec['sheet']} r.{rec['row']}",
                     f"{rec['status_raw']} <-> {names}"))
    return {}


def build_teachers(records, resolve):
    """Groups the subject rows by teacher. Returns (teachers, warnings)."""
    teachers, warnings = OrderedDict(), []
    spellings = defaultdict(Counter)

    for rec in records:
        people = [(p, 'lead') for p in rec['leads']] + [(p, 'co') for p in rec['cos']]
        if not people:
            warnings.append(('Õppejõuta rida (ei lähe ühegi faili)', f"{rec['sheet']} r.{rec['row']}",
                             f"{rec['group_primary']} | {rec['code']} | {rec['name']}"))
            continue
        row_status = assign_row_status(rec, people, warnings)
        seen = set()
        for (pname, marker, note), role in people:
            key = resolve(pname)
            if key in seen:  # the same person twice in one row
                continue
            seen.add(key)
            t = teachers.setdefault(key, dict(rows=[], votes=Counter(), tokens=set(), sheets=Counter()))
            spellings[key][clean(pname)] += 1
            if marker:
                t['votes'][marker] += 2  # 'TT' written next to the name is the strongest sign
            if pname in row_status:
                t['votes'][row_status[pname]] += 1
            t['tokens'].update(rec['tokens'])
            t['sheets'][rec['sheet']] += 1
            t['rows'].append(dict(rec=rec, role=role, note=note, key=key))

    for key, t in teachers.items():
        names = sorted(spellings[key].items(), key=lambda kv: (-kv[1], -len(kv[0])))
        t['name'] = names[0][0]
        if len(names) > 1:
            warnings.append(('Nime variandid (liidetud ühte faili)', t['name'], ' / '.join(n for n, _ in names)))
        top = t['votes'].most_common()
        if not top:
            t['status'] = None
            warnings.append(('Staatus puudub (TT / lepinguline)', t['name'], ''))
        elif len(top) > 1 and top[0][1] == top[1][1]:
            t['status'] = None
            warnings.append(('Staatus vastuoluline', t['name'], repr(dict(t['votes']))))
        else:
            t['status'] = top[0][0]

    for rec in records:
        for text in rec['free_notes']:
            warnings.append(('Tekst õppejõu lahtris (lisatud märkusena)', f"{rec['sheet']} r.{rec['row']}", text))
    keys = list(teachers)
    for i, a in enumerate(keys):
        for b in keys[i + 1:]:
            if SequenceMatcher(None, a, b).ratio() >= 0.85:
                warnings.append(('Võimalik sama inimene (vajadusel NAME_OVERRIDES)',
                                 teachers[a]['name'], teachers[b]['name']))
    return teachers, warnings


# ---------------------------------------------------------------------------
# Block 2: colour rules (read from the legend of the calendar)
# ---------------------------------------------------------------------------
def read_legend(cal, modes):
    """Legend = row 2 of the calendar sheet. For every coloured legend cell:
       ('groups', {(programme, year)...})  - session colour of some groups,
       ('always', None)                    - holidays, general meeting ... (everybody)."""
    codes = '|'.join(sorted(modes, key=len, reverse=True))
    pattern = re.compile(r'(' + codes + r')\s*(\d(?:\s*,\s*\d(?![A-Za-zÕÄÖÜõäöü]))*)')
    rules = OrderedDict()
    for cell in cal[2]:
        text = clean(cell.value)
        if not text or not cell.fill.fill_type or cell.fill.fgColor.type != 'rgb':
            continue
        rgb = cell.fill.fgColor.rgb[-6:]
        tokens = {(m.group(1), int(y)) for m in pattern.finditer(text) for y in re.findall(r'\d', m.group(2))}
        if tokens:
            rules[rgb] = ('groups', tokens)
        else:
            rules[rgb] = ('always', None)
    return rules


def rule_applies(rule, teacher_tokens, modes):
    """Does this colour belong to the teacher (stay coloured) or not (becomes grey)?"""
    kind, tokens = rule
    if kind == 'always':
        return True
    return any(p == lp and (y is None or y == ly) for p, y in teacher_tokens for lp, ly in tokens)


def row_color(subject_tokens, rules):
    """Colour of a subject row = the session colour of its group (if it has one)."""
    for rgb, (kind, tokens) in rules.items():
        if kind == 'groups' and any(p == lp and (y is None or y == ly)
                                    for p, y in subject_tokens for lp, ly in tokens):
            return rgb
    return None


# ---------------------------------------------------------------------------
# Block 3: dates of the calendar
# ---------------------------------------------------------------------------
def calendar_dates(cal):
    """Returns {(row, col): correct date}.
    The calendar is a grid: column = week (Monday + 7 days), date row = weekday.
    First Monday = cell of the first week; the YEAR is taken from the sheet name
    (the stored years in the table are not reliable, only 'day-month' is shown)."""
    day_rows = {}
    for r in range(1, cal.max_row + 1):
        label = clean(cal.cell(r, 1).value)
        if label in DAY_LABELS and label not in day_rows:
            day_rows[label] = r - 1  # the date row is right above the day block
    first_cell = cal.cell(day_rows['E'], 3).value
    year_in_title = re.search(r'(\d{4})', cal.title)
    if not isinstance(first_cell, (datetime, date)) or not year_in_title:
        return {}
    start = date(int(year_in_title.group(1)), first_cell.month, first_cell.day)
    if start.weekday() != 0:  # must be a Monday, otherwise do not touch
        return {}
    result = {}
    for label, r in day_rows.items():
        for col in range(3, cal.max_column + 1):
            if isinstance(cal.cell(r, col).value, (datetime, date)):
                result[(r, col)] = start + timedelta(days=7 * (col - 3) + DAY_LABELS.index(label))
    return result


# ---------------------------------------------------------------------------
# Block 4: building one teacher file
# ---------------------------------------------------------------------------
def group_display(rec):
    other = rec['group_other']
    if other and other != rec['group_primary'] and not other.casefold().startswith(PLACEHOLDERS) and '?' not in other:
        return f"{rec['group_primary']} + {other}"
    return rec['group_primary']


def info_text(item, resolve):
    """Text for the column 'Lisainfo': everything useful the shared table says about the subject."""
    rec, role = item['rec'], item['role']
    parts = []
    co_names = ', '.join(p[0] + (f' ({p[2]})' if p[2] else '') for p in rec['cos'])
    if role == 'co':
        lead = ', '.join(p[0] for p in rec['leads'])
        parts.append('Kaasõppejõud' + (f', vastutav õppejõud: {lead}' if lead else ''))
        others = [p[0] for p in rec['cos'] if resolve(p[0]) != item['key']]
        if others:
            parts.append('Koos: ' + ', '.join(others))
    else:
        if item['note']:
            parts.append(f"Märkus: {item['note']}")
        if co_names:
            parts.append(f'Koos: {co_names}')
    if rec['subgrp']:
        parts.append(f"Rühmad õppeaines: {rec['subgrp']}")
    if rec['ngroups']:
        parts.append(f"Rühmi: {rec['ngroups']}")
    if rec['rgroup']:
        parts.append(f"Rühm: {rec['rgroup']}")
    if rec['type'] and not normalize(rec['type']).startswith('kohustuslik'):
        parts.append(rec['type'])
    if rec['module'] and rec['module'] != rec['group_primary']:
        parts.append(f"Moodul: {rec['module']}")
    if rec['notes']:
        parts.append(rec['notes'])
    parts.extend(rec['free_notes'])
    c = rec['contact']  # in some sheets this column holds comments
    if c and '@' not in c and not re.search(r'\d{5,}', c):
        parts.append(c)
    return '; '.join(parts)


def status_title(status):
    return {'TT': 'TT (külalisõppejõud)', 'lepinguline': 'lepinguline (kolledži õppejõud)'}.get(
        status, 'staatus täpsustamata')


def build_teacher_file(master_path, key, t, resolve, semester):
    wb = openpyxl.load_workbook(master_path)  # reload: keeps the colour theme of the table
    cal, abi, _ = find_special_sheets(wb)
    modes = read_study_modes(abi)
    rules = read_legend(cal, modes)
    new_dates = calendar_dates(cal)

    rows = t['rows']
    n = len(rows)
    offset = n + 1  # calendar row 2 (legend) -> row n+3
    ws = wb.create_sheet(f'{semester} Kevadsemestri põhi')

    # ---- title ---------------------------------------------------------------
    ws['A1'] = f"{t['name']} — {status_title(t['status'])}"
    ws['A1'].font = Font(name=FONT, size=14, bold=True)
    ws.row_dimensions[1].height = 24

    # ---- subject table (like the sample: header row 2, subjects from row 3) ----
    thin = Side(style='thin', color='000000')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left = Alignment(horizontal='left', vertical='center', wrap_text=True)
    headers = {
        1: 'Rühmad', 2: 'Õppeaine kood', 3: 'Õppeaine nimetus', 9: 'EAP',
        10: 'Soovituslik paaride arv (90 min)',
        11: 'Kontakttundide soov 90 min (kui soovite vähem/rohkem kontakttunde tunniplaani, '
            'siis palun siia täpset arvu)',
        12: 'Mitu paari järjest võib panna?',
        13: 'Kas aine lõpeb eksamiga (E) või arvestusega (A)?',
        14: 'Eksami ajad', 15: 'Arvestuse aeg', 16: 'Soovitud auditoorium',
        17: 'Lisainfo \ne-õpe, väljasõit, ettevõtete külastus, külalisõppejõud. Kui õppeainel on mitu '
            'õppejõudu, siis palun märkida siia kuidas on õppetöö korraldatud jne',
    }
    header_size = {11: 8, 12: 9, 13: 8, 14: 9, 15: 9, 16: 8}  # font sizes of the sample
    for col in range(1, 23):
        c = ws.cell(2, col, headers.get(col))
        c.font = Font(name=FONT, size=header_size.get(col, 11), bold=True)
        c.fill, c.alignment, c.border = fill(HEADER_BLUE), center, border
    ws.row_dimensions[2].height = 93.6

    for i, item in enumerate(rows):
        r, rec = 3 + i, item['rec']
        info = info_text(item, resolve)
        group_text = group_display(rec)
        values = {1: group_text, 2: rec['code'], 3: rec['name'], 9: rec['eap'], 10: rec['pairs'], 17: info}
        color = row_color(rec['tokens'], rules)
        for col in range(1, 23):
            c = ws.cell(r, col, values.get(col))
            c.font, c.border = Font(name=FONT, size=11), border
            c.alignment = left if col in (1, 3, 17) else center
            if color and col <= 11:  # the sample colours columns A..K
                c.fill = fill(color)
        lines = max(math.ceil(len(rec['name']) / 70), math.ceil(len(info) / 70),
                    math.ceil(len(group_text) / 10), math.ceil(len(rec['code']) / 14), 1)
        ws.row_dimensions[r].height = max(21.75, 15 * lines + 4)
    for r in range(2, 3 + n):
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=8)  # name C:H
        ws.merge_cells(start_row=r, start_column=17, end_row=r, end_column=22)  # info Q:V

    # ---- general calendar (values, styles, merges, sizes) -----------------------
    for row in cal.iter_rows(min_row=2, max_row=cal.max_row):
        for c in row:
            d = ws.cell(c.row + offset, c.column, c.value)
            if c.has_style:
                d._style = copy(c._style)
    for (r, col), day in new_dates.items():  # correct dates (same cell style)
        d = ws.cell(r + offset, col)
        d.value = datetime(day.year, day.month, day.day)
        d.number_format = 'd-mmm'
    for mr in cal.merged_cells.ranges:
        if mr.min_row >= 2:
            ws.merge_cells(start_row=mr.min_row + offset, end_row=mr.max_row + offset,
                           start_column=mr.min_col, end_column=mr.max_col)
    for r, dim in cal.row_dimensions.items():
        if r >= 2 and dim.height:
            ws.row_dimensions[r + offset].height = dim.height
    for letter, dim in cal.column_dimensions.items():
        nd = ws.column_dimensions[letter]
        nd.min, nd.max, nd.width, nd.hidden = dim.min, dim.max, dim.width, dim.hidden
    if 'A' not in cal.column_dimensions:  # column A holds the group names in the top table
        ws.column_dimensions['A'].width = 13.7  # (width of the sample)

    # ---- legend: add "Lühendatud tööpäev" (used in the calendar, missing in the legend) ----
    legend_row = 2 + offset
    exam = next(c for c in ws[legend_row] if 'eksam' in normalize(c.value))
    col = exam.column + 3
    for k in (0, 1):
        cell = ws.cell(legend_row, col + k)
        cell._style = copy(exam._style)
        cell.fill = fill(SHORT_DAY)
    ws.cell(legend_row, col).value = 'Lühendatud tööpäev'
    ws.merge_cells(start_row=legend_row, start_column=col, end_row=legend_row, end_column=col + 1)
    ws.row_dimensions[legend_row].height = 72

    # ---- optional deadlines block (as in the sample) ----
    if DEADLINES_TEXT:
        ws.merge_cells(start_row=2, start_column=26, end_row=legend_row, end_column=33)
        c = ws.cell(2, 26, DEADLINES_TEXT)
        c.font, c.alignment = Font(name=FONT, size=11, bold=True), Alignment(wrap_text=True, vertical='top')

    # ---- colours: the legend keeps all colours; in the calendar only the teacher's groups ----
    last_row = cal.max_row + offset
    for row in ws.iter_rows(min_row=legend_row + 1, max_row=last_row):
        for c in row:
            if not c.fill.fill_type or c.fill.fgColor.type != 'rgb':
                continue
            rule = rules.get(c.fill.fgColor.rgb[-6:])
            if rule and not rule_applies(rule, t['tokens'], modes):
                c.fill = fill(GREY)

    # ---- bottom table "Soovituslikud kontakttundide arvud" (Abiinfo) ----
    top = next(r for r in range(1, abi.max_row + 1) if normalize(abi.cell(r, 2).value).startswith('soovituslikud'))
    start = last_row + 1
    for k in range(6):
        for col in range(2, 7):
            s = abi.cell(top + k, col)
            d = ws.cell(start + k, col - 1, s.value)
            if s.has_style:
                d._style = copy(s._style)
    ws.merge_cells(start_row=start + 5, start_column=2, end_row=start + 5, end_column=5)

    # ---- keep only the new sheet, save ----
    for name in wb.sheetnames:
        if name != ws.title:
            del wb[name]
    ws.sheet_view.zoomScale = 80

    suffix = ' TT' if t['status'] == 'TT' else ''
    path = os.path.join(OUTPUT_DIR, safe_filename(f"{t['name']}{suffix}") + f'_Kevadsemester{semester}.xlsx')
    try:
        wb.save(path)
        return path, len(new_dates)
    except PermissionError:
        print(f"  !! Cannot save '{path}'. Close the file in Excel and run again.")
        return None, 0


# ---------------------------------------------------------------------------
# Block 5: report
# ---------------------------------------------------------------------------
def write_report(teachers, warnings, files):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Õppejõud'
    ws.append(['Õppejõud', 'Staatus', 'Ainete arv', 'Erialalehed', 'Rühmad', 'Fail'])
    for key, t in sorted(teachers.items(), key=lambda kv: kv[1]['name'].casefold()):
        groups = ', '.join(sorted(f'{p} {y}' if y else p for p, y in t['tokens']))
        ws.append([t['name'], t['status'] or '?', len(t['rows']), ', '.join(t['sheets']), groups,
                   os.path.basename(files.get(key) or '')])
    w2 = wb.create_sheet('Hoiatused')
    w2.append(['Tüüp', 'Kus / kes', 'Detail'])
    for w in warnings:
        w2.append(list(w))
    for sheet in (ws, w2):
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = min(
                70, max(12, max(len(str(c.value or '')) for c in col) + 2))
        for c in sheet[1]:
            c.font = Font(bold=True)
    wb.save(os.path.join(OUTPUT_DIR, '_aruanne.xlsx'))


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
def main():
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    wb = openpyxl.load_workbook(INPUT_FILE)
    cal, abi, program_sheets = find_special_sheets(wb)
    modes = read_study_modes(abi)
    m = re.search(r'(\d{4})\s*/\s*(\d{2})', clean(cal['A3'].value))
    semester = f'{m.group(1)}_{m.group(2)}' if m else 'semester'

    records = []
    for sheet in program_sheets:
        records += read_program_sheet(sheet, set(modes))
    resolve = NameResolver([p[0] for r in records for p in r['leads'] + r['cos']])
    teachers, warnings = build_teachers(records, resolve)
    print(f'Specialty sheets: {len(program_sheets)}, subject rows: {len(records)}, teachers: {len(teachers)}')

    files, fixed_dates = {}, 0
    for key, t in teachers.items():
        path, fixed_dates = build_teacher_file(INPUT_FILE, key, t, resolve, semester)
        files[key] = path
        print(f"-> {t['name']:<28} {str(t['status']):<12} subjects: {len(t['rows'])}")
    assert len({p for p in files.values() if p}) == len([p for p in files.values() if p]), 'duplicate file names!'

    write_report(teachers, warnings, files)
    print(f'\nDone: {sum(1 for p in files.values() if p)} files in output_files/, '
          f'warnings: {len(warnings)} (see _aruanne.xlsx)')


if __name__ == '__main__':
    main()